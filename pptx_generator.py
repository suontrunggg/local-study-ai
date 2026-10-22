#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
PPTX Generator — Tạo bài thuyết trình PowerPoint (.pptx) chuẩn Swiss Design & PPT-Master
Hỗ trợ:
  - Tự động bóc tách dàn ý Slide & Speaker Notes từ phản hồi AI
  - Thiết kế slide 16:9 Widescreen với bảng màu Thụy Sĩ (Swiss Design)
  - Khung nội dung tối giản, thẻ Eyebrow, bố cục thẻ nổi
  - Tích hợp ghi chú diễn giả (Speaker Notes) trực tiếp vào file .pptx
"""
import os
import re
import uuid
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def parse_slides_from_markdown(text: str):
    """
    Tự động phân tích văn bản Markdown từ AI để trích xuất danh sách các slide.
    Hỗ trợ định dạng chuẩn:
      ### 📑 [SLIDE 1]: Tiêu đề
      - **🎯 Thông điệp cốt lõi:** ...
      - **📊 Bố cục nội dung chính:** ...
      - **💡 Ghi chú thuyết trình:** ...
    Kèm 3 tầng fallback thông minh để không bao giờ bỏ sót slide.
    """
    slides = []
    # Tìm tiêu đề chung của bài báo cáo nếu có
    main_title_match = re.search(r"^#\s+(.+)$", text, re.MULTILINE)
    main_title = main_title_match.group(1).strip() if main_title_match else "Báo Cáo Thuyết Trình Nghiên Cứu"
    main_title = re.sub(r"[\*#_`]", "", main_title).strip()

    # Pass 1: Định dạng chuẩn ### [SLIDE 1] hoặc ### Slide 1 hoặc ## Slide 1 hoặc **Slide 1:**
    pattern_standard = r"(?=(?:###|##|\*\*)\s*(?:📑\s*)?(?:\[\s*)?(?:SLIDE|Slide|Trang|Phần)\s*\d+)"
    chunks = re.split(pattern_standard, text, flags=re.IGNORECASE)

    for chunk in chunks:
        chunk = chunk.strip()
        if not chunk:
            continue
        header_match = re.search(r"(?:###|##|\*\*)\s*(?:📑\s*)?(?:\[\s*)?(?:SLIDE|Slide|Trang|Phần)\s*(\d+)\]?[:\s\*\-]+(.+)", chunk, re.IGNORECASE)
        if not header_match:
            continue
        
        slide_num = header_match.group(1)
        raw_title = re.sub(r"[\*#_`]", "", header_match.group(2)).strip()
        raw_title = raw_title.split("\n")[0].strip()

        # Trích xuất Speaker Notes
        speaker_notes = ""
        notes_match = re.search(r"(?:Ghi chú thuyết trình|Speaker Notes|Ghi chú)[:\s\*\-]+([^\n]+(?:\n(?!\s*[-*#]).+)*)", chunk, re.IGNORECASE)
        if notes_match:
            speaker_notes = re.sub(r"[\*`]", "", notes_match.group(1)).strip()

        # Trích xuất Thông điệp cốt lõi
        takeaway = ""
        takeaway_match = re.search(r"(?:Thông điệp cốt lõi|Key Takeaway|Mục tiêu)[:\s\*\-]+([^\n]+)", chunk, re.IGNORECASE)
        if takeaway_match:
            takeaway = re.sub(r"[\*`]", "", takeaway_match.group(1)).strip()

        # Trích xuất các ý chính (bullets)
        bullets = []
        if takeaway:
            bullets.append(f"🎯 Điểm cốt lõi: {takeaway}")

        lines = chunk.split("\n")
        for line in lines:
            line_str = line.strip()
            if re.match(r"^[\*\-•]\s+", line_str):
                cleaned = re.sub(r"^[\*\-•]\s+", "", line_str)
                if any(k in cleaned for k in ["Thông điệp cốt lõi", "Bố cục nội dung chính", "Ghi chú thuyết trình", "Key Takeaway", "Speaker Notes"]):
                    continue
                cleaned = re.sub(r"[\*`]", "", cleaned).strip()
                if cleaned and cleaned not in bullets:
                    bullets.append(cleaned)

        slides.append({
            "slide_num": slide_num,
            "title": raw_title or f"Slide {slide_num}",
            "category": "PHÂN TÍCH HỌC THUẬT",
            "bullets": bullets[:5],
            "speaker_notes": speaker_notes
        })

    # Pass 2: Fallback nếu AI dùng ### hoặc ## nhưng không ghi chữ "Slide"
    if len(slides) < 2:
        h_chunks = re.split(r"(?=###\s+)", text)
        if len(h_chunks) < 2:
            h_chunks = re.split(r"(?=##\s+)", text)

        if len(h_chunks) >= 2:
            slides = []
            for idx, chunk in enumerate(h_chunks):
                chunk = chunk.strip()
                if not chunk:
                    continue
                h_match = re.search(r"^(?:###|##)\s+(.+)", chunk, re.MULTILINE)
                if not h_match:
                    continue
                slide_title = re.sub(r"[\*#_`]", "", h_match.group(1)).split("\n")[0].strip()
                if len(slide_title) < 3:
                    continue
                
                bullets = []
                lines = chunk.split("\n")
                for line in lines:
                    line_str = line.strip()
                    if re.match(r"^[\*\-•\d+\.]\s+", line_str):
                        cleaned = re.sub(r"^[\*\-•\d+\.]\s+", "", line_str)
                        cleaned = re.sub(r"[\*`]", "", cleaned).strip()
                        if cleaned and len(cleaned) > 5 and cleaned not in bullets:
                            bullets.append(cleaned)

                slides.append({
                    "slide_num": str(len(slides) + 1),
                    "title": slide_title,
                    "category": "BÁO CÁO HỌC THUẬT",
                    "bullets": bullets[:5] if bullets else ["Nội dung phân tích trọng tâm từ tài liệu", "Số liệu và luận điểm thực nghiệm"],
                    "speaker_notes": f"Trình bày chi tiết về {slide_title}."
                })

    # Pass 3: Fallback phân đoạn văn bản nếu AI chỉ viết văn xuôi
    if len(slides) < 2 and len(text.strip()) > 150:
        paragraphs = [p.strip() for p in text.split("\n\n") if len(p.strip()) > 40]
        if paragraphs:
            slides = []
            for i, p in enumerate(paragraphs[:6]):
                first_sentence = p.split(".")[0].strip()
                slide_title = first_sentence[:50] if len(first_sentence) > 10 else f"Nội dung trọng tâm {i+1}"
                slides.append({
                    "slide_num": str(i + 1),
                    "title": slide_title,
                    "category": "TỔNG HỢP NỘI DUNG",
                    "bullets": [p[:120] + "..."] if len(p) > 120 else [p],
                    "speaker_notes": "Diễn giải chi tiết nội dung slide cho hội đồng."
                })

    return main_title, slides


def create_swiss_deck(title: str, slides: list, output_path: str):
    """
    Sinh file .pptx chuẩn 16:9 phong cách Swiss Design cao cấp.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # --- SLIDE 1: COVER (NỀN TỐI SANG TRỌNG) ---
    slide0 = prs.slides.add_slide(blank_layout)
    bg0 = slide0.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(7.5))
    bg0.fill.solid()
    bg0.fill.fore_color.rgb = RGBColor(15, 23, 42) # Midnight Slate #0F172A
    bg0.line.fill.background()

    # Khung tiêu đề bìa
    tb0 = slide0.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(11.0), Inches(3.2))
    tf0 = tb0.text_frame
    tf0.word_wrap = True
    p0 = tf0.paragraphs[0]
    p0.text = title.upper()
    p0.font.size = Pt(38)
    p0.font.bold = True
    p0.font.color.rgb = RGBColor(255, 255, 255)

    p_sub = tf0.add_paragraph()
    p_sub.text = "BÁO CÁO NGHIÊN CỨU & BẢO VỆ ĐỀ TÀI • SWISS DESIGN EDITION"
    p_sub.font.size = Pt(16)
    p_sub.font.bold = True
    p_sub.font.color.rgb = RGBColor(56, 189, 248) # Sky 400
    p_sub.space_before = Pt(24)

    p_author = tf0.add_paragraph()
    p_author.text = "Tạo bởi Study AI Local • Tự động cấu trúc luận điểm & Speaker Notes"
    p_author.font.size = Pt(13)
    p_author.font.color.rgb = RGBColor(148, 163, 184) # Slate 400
    p_author.space_before = Pt(12)

    # --- CÁC SLIDE NỘI DUNG (NỀN SÁNG TỐI GIẢN) ---
    for i, s in enumerate(slides, 1):
        slide = prs.slides.add_slide(blank_layout)

        # Nền sáng
        bg_s = slide.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(7.5))
        bg_s.fill.solid()
        bg_s.fill.fore_color.rgb = RGBColor(248, 250, 252) # Slate 50
        bg_s.line.fill.background()

        # Eyebrow tag
        tb_eye = slide.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.3), Inches(0.4))
        p_eye = tb_eye.text_frame.paragraphs[0]
        p_eye.text = f"SLIDE {i:02d} • {s.get('category', 'PHÂN TÍCH')}"
        p_eye.font.size = Pt(12)
        p_eye.font.bold = True
        p_eye.font.color.rgb = RGBColor(0, 113, 227) # Apple Blue

        # Tiêu đề slide
        tb_t = slide.shapes.add_textbox(Inches(1.0), Inches(0.95), Inches(11.3), Inches(0.9))
        tb_t.text_frame.word_wrap = True
        p_t = tb_t.text_frame.paragraphs[0]
        p_t.text = s.get("title", f"Slide {i}")
        p_t.font.size = Pt(25)
        p_t.font.bold = True
        p_t.font.color.rgb = RGBColor(15, 23, 42)

        # Thẻ nội dung (Card Box)
        card = slide.shapes.add_shape(1, Inches(1.0), Inches(2.0), Inches(11.333), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(255, 255, 255)
        card.line.color.rgb = RGBColor(226, 232, 240)
        card.line.width = Pt(1)

        tf_card = card.text_frame
        tf_card.word_wrap = True
        tf_card.margin_left = Inches(0.5)
        tf_card.margin_right = Inches(0.5)
        tf_card.margin_top = Inches(0.5)

        bullets = s.get("bullets", [])
        if not bullets:
            bullets = ["Nội dung slide đang được cập nhật."]

        first = True
        for bullet in bullets:
            if first:
                p_b = tf_card.paragraphs[0]
                first = False
            else:
                p_b = tf_card.add_paragraph()
            p_b.text = "•  " + bullet
            p_b.font.size = Pt(16)
            p_b.font.color.rgb = RGBColor(51, 65, 85) # Slate 700
            p_b.space_after = Pt(16)

        # Ghi chú thuyết trình (Speaker Notes)
        notes = s.get("speaker_notes", "")
        if notes:
            slide.notes_slide.notes_text_frame.text = f"[GHI CHÚ HỘI ĐỒNG]:\n{notes}"

    # --- SLIDE CUỐI: KẾT LUẬN & HỎI ĐÁP (Q&A) ---
    slide_end = prs.slides.add_slide(blank_layout)
    bg_end = slide_end.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(7.5))
    bg_end.fill.solid()
    bg_end.fill.fore_color.rgb = RGBColor(15, 23, 42)
    bg_end.line.fill.background()

    tb_end = slide_end.shapes.add_textbox(Inches(1.2), Inches(2.4), Inches(11.0), Inches(2.8))
    tf_end = tb_end.text_frame
    tf_end.word_wrap = True
    p_end0 = tf_end.paragraphs[0]
    p_end0.text = "CẢM ƠN QUÝ THẦY CÔ & HỘI ĐỒNG"
    p_end0.font.size = Pt(36)
    p_end0.font.bold = True
    p_end0.font.color.rgb = RGBColor(255, 255, 255)

    p_end1 = tf_end.add_paragraph()
    p_end1.text = "Sẵn sàng cho phần Hỏi & Đáp (Q&A) • Trợ lý Nghiên cứu Study AI"
    p_end1.font.size = Pt(17)
    p_end1.font.color.rgb = RGBColor(56, 189, 248)
    p_end1.space_before = Pt(20)

    prs.save(output_path)
    return output_path
